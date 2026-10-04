import io

import pytest
from openpyxl import Workbook

from app.datasets import DatasetError, QueryError, load_upload


def test_samples_load_with_expected_tables(registry):
    assert set(registry.samples) == {"store", "food", "hr"}
    store = registry.get("store")
    assert {t.name: t.row_count for t in store.tables} == {"orders": 20000, "customers": 3000, "products": 120}
    orders = next(t for t in store.tables if t.name == "orders")
    amount = next(c for c in orders.columns if c.name == "amount")
    assert amount.kind == "number" and amount.description


def test_sample_data_is_deterministic(registry):
    total = registry.get("store").execute("SELECT ROUND(SUM(amount), 2) FROM orders").rows[0][0]
    again = registry.get("store").execute("SELECT ROUND(SUM(amount), 2) FROM orders").rows[0][0]
    assert total == again and total > 0


def test_schema_text_lists_columns_and_examples(registry):
    text = registry.get("hr").schema_text()
    assert "TABLE employees" in text and "department_id INTEGER" in text and "e.g." in text


def test_file_access_is_blocked_inside_the_database(registry):
    with pytest.raises(QueryError, match="disabled"):
        registry.get("store").execute("SELECT * FROM read_csv('/etc/passwd')")


def test_configuration_cannot_be_changed(registry):
    with pytest.raises(QueryError):
        registry.get("store").execute("SET enable_external_access = true")


def test_slow_queries_are_stopped(registry):
    with pytest.raises(QueryError, match="longer than"):
        registry.get("store").execute(
            "SELECT SUM(a.range * b.range) FROM range(200000) a, range(200000) b", timeout_s=0.3)


def test_results_are_capped_and_json_safe(registry):
    r = registry.get("store").execute("SELECT order_id, order_date, amount FROM orders", max_rows=10)
    assert len(r.rows) == 10 and r.truncated
    assert isinstance(r.rows[0][1], str) and isinstance(r.rows[0][2], float)


def test_csv_upload_is_profiled():
    csv = b"Region,Month,Sales\nNorth,2026-01-01,100\nSouth,2026-01-01,250\nNorth,2026-02-01,130\n"
    ds = load_upload("My Sales.csv", csv)
    assert ds.kind == "upload" and ds.tables[0].name == "my_sales"
    cols = {c.name: c.kind for c in ds.tables[0].columns}
    assert cols == {"Region": "text", "Month": "date", "Sales": "number"}
    assert ds.execute('SELECT SUM("Sales") FROM my_sales').rows[0][0] == 480


def test_xlsx_upload_turns_each_sheet_into_a_table():
    wb = Workbook()
    ws = wb.active
    ws.title = "Staff"
    ws.append(["Name", "Team", "Salary"])
    ws.append(["Asha", "Ops", 50000])
    ws.append(["Ravi", "Tech", 70000])
    ws2 = wb.create_sheet("Teams")
    ws2.append(["Team", "Floor"])
    ws2.append(["Ops", 2])
    buf = io.BytesIO()
    wb.save(buf)
    ds = load_upload("company.xlsx", buf.getvalue())
    assert [t.name for t in ds.tables] == ["staff", "teams"]
    assert ds.execute("SELECT MAX(Salary) FROM staff").rows[0][0] == 70000


@pytest.mark.parametrize("name, data, message", [
    ("notes.txt", b"hello", ".csv or .xlsx"),
    ("empty.csv", b"   ", "empty"),
    ("broken.xlsx", b"not really excel", "Could not read"),
    ("header_only.csv", b"a,b,c\n", "no rows"),
])
def test_bad_uploads_give_clear_errors(name, data, message):
    with pytest.raises(DatasetError, match=message):
        load_upload(name, data)


def test_upload_row_limit():
    data = b"x\n" + b"\n".join(str(i).encode() for i in range(50)) + b"\n"
    with pytest.raises(DatasetError, match="limit"):
        load_upload("big.csv", data, max_rows=10)
